#!/usr/bin/env python3
"""Test end-to-end di scripts/fatturapa/parse_fatture.py su fixture sintetiche.

Esegue: python tests/test_parse_fatture.py
Copre: XML semplice, XML.P7M (firma CAdES generata con openssl), ZIP annidato, nota di credito TD04
(segno negativo), filtro --anno, classificazione e flag "da verificare", quadratura.
"""
import shutil
import subprocess
import sys
import tempfile
import zipfile
from decimal import Decimal
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts" / "fatturapa"))
import parse_fatture as pf  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

FIX = ROOT / "tests" / "fixtures"


def make_p7m(src: Path, dst: Path) -> bool:
    """Firma CAdES con certificato self-signed usa-e-getta (come un vero .xml.p7m)."""
    with tempfile.TemporaryDirectory() as td:
        key, crt = Path(td) / "k.pem", Path(td) / "c.pem"
        r = subprocess.run(["openssl", "req", "-x509", "-newkey", "rsa:2048", "-nodes", "-keyout", key,
                            "-out", crt, "-days", "1", "-subj", "/CN=test"], capture_output=True)
        if r.returncode != 0:
            return False
        r = subprocess.run(["openssl", "cms", "-sign", "-binary", "-nodetach", "-outform", "DER", "-in", src,
                            "-signer", crt, "-inkey", key, "-out", dst], capture_output=True)
        return r.returncode == 0


def main():
    with tempfile.TemporaryDirectory() as td:
        work = Path(td)
        (work / "in").mkdir()
        shutil.copy(FIX / "IT01274580248_00001.xml", work / "in")
        shutil.copy(FIX / "IT01274580248_00099_2024.xml", work / "in")
        p7m_ok = make_p7m(FIX / "IT01234567890_00055.xml", work / "in" / "IT01234567890_00055.xml.p7m")
        if not p7m_ok:  # openssl assente: usa l'xml in chiaro
            shutil.copy(FIX / "IT01234567890_00055.xml", work / "in")
        with zipfile.ZipFile(work / "in" / "lotto.zip", "w") as z:
            z.write(FIX / "IT01274580248_00002_NC.xml", "sub/IT01274580248_00002_NC.xml")
        out = work / "out" / "acquisti.xlsx"
        rc = pf.main(["--in", str(work / "in"), "--out", str(out), "--anno", "2025", "--piva", "01241890258", "--csv"])
        assert rc == 0, "exit code"
        wb = load_workbook(out, data_only=True)
        isa = {r[0]: r for r in wb["Totali_ISA"].iter_rows(min_row=2, values_only=True)}
        exp = {  # confermato, da verificare, totale
            "carne": (Decimal("122.50"), Decimal("0"), Decimal("122.50")),
            "sfarinati": (Decimal("101.60"), Decimal("0"), Decimal("101.60")),
            "vino": (Decimal("108.00"), Decimal("75.00"), Decimal("183.00")),
            "birra": (Decimal("85.00"), Decimal("0"), Decimal("85.00")),
        }
        for cat, (ok, ver, tot) in exp.items():
            got = isa[cat]
            assert Decimal(str(got[1])) == ok, f"{cat} confermato {got[1]} != {ok}"
            assert Decimal(str(got[2])) == ver, f"{cat} da verificare {got[2]} != {ver}"
            assert Decimal(str(got[3])) == tot, f"{cat} totale {got[3]} != {tot}"
        righe = list(wb["Righe"].iter_rows(min_row=2, values_only=True))
        cats = {r[7]: (r[14] or "", r[15] or "", r[16] or "") for r in righe}
        assert cats["TONNO PINNE GIALLE FILETTI"][0] == "pesce"
        assert cats["SPESE DI TRASPORTO"][0] == "accessoria"
        assert cats["DETERSIVO LAVASTOVIGLIE 12KG"][0] == "non_alimentare"
        assert cats["ACQUA NATURALE VETRO 1L"][0] == "bevande_analcoliche"
        assert cats["SPECK ALTO ADIGE IGP META'"] == ("carne", "salumi", "")
        assert cats["BIRRA FORST KRONEN FUSTO 30 LT"][1] == "fusti"
        assert cats["VINO BIANCO SOAVE DOC 75CL"][2] == "SI", "aliquota 10% su vino deve essere segnalata"
        assert cats["ART. 4471 CONF. 6 BT"] == ("vino", "", "SI"), "default fornitore Schinosa"
        assert "CARNE DI CERVO SPEZZATINO" not in cats, "fattura 2024 esclusa dal filtro --anno"
        fatt = list(wb["Fatture"].iter_rows(min_row=2, values_only=True))
        assert len(fatt) == 3, f"attese 3 fatture 2025, trovate {len(fatt)}"
        nc = [f for f in fatt if f[3] == "TD04"][0]
        assert Decimal(str(nc[7])) == Decimal("-10.80"), "nota di credito negativa"
        assert list(wb["Quadratura"].iter_rows(min_row=2, values_only=True)) == [], "nessuno scostamento atteso"
        assert out.with_suffix(".righe.csv").exists()
        print(f"OK: 3 fatture (p7m={'sì' if p7m_ok else 'no'}), {len(righe)} righe, totali ISA corretti")


if __name__ == "__main__":
    main()
