#!/usr/bin/env python3
"""Test offline di scripts/stripe/export_stripe.py con la fixture stripe_demo.json."""
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts" / "stripe"))
import export_stripe as es  # noqa: E402
from openpyxl import load_workbook  # noqa: E402


def main():
    with tempfile.TemporaryDirectory() as td:
        out = Path(td) / "stripe.xlsx"
        rc = es.main(["--demo", str(ROOT / "tests/fixtures/stripe_demo.json"), "--from", "2025-01-01", "--to", "2025-12-31", "--out", str(out)])
        assert rc == 0
        wb = load_workbook(out, data_only=True)
        rows = {r[0]: r for r in wb["Mensile"].iter_rows(min_row=2, values_only=True)}
        jan, feb, tot = rows["2025-01"], rows["2025-02"], rows["TOTALE"]
        assert (jan[1], jan[2], jan[3], jan[6], jan[7]) == (430.0, -50.0, 11.3, 368.7, 368.7), jan
        assert (feb[1], feb[3], feb[6], feb[7]) == (300.0, 6.0, 294.0, 0.0), feb
        assert (tot[1], tot[3], tot[7]) == (730.0, 17.3, 368.7), tot
        assert len(list(wb["Transazioni"].iter_rows(min_row=2))) == 6
        assert len(list(wb["Payout"].iter_rows(min_row=2))) == 1
        q = {r[0]: r[1] for r in wb["Quadratura"].iter_rows(values_only=True) if r and r[0]}
        assert abs(q["Differenza = variazione del saldo Stripe nel periodo"] - 294.0) < 0.001
        print("OK: export Stripe demo, totali mensili e quadratura corretti")


if __name__ == "__main__":
    main()
