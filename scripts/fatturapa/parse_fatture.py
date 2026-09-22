#!/usr/bin/env python3
"""
parse_fatture.py – Dalle fatture elettroniche ricevute (FatturaPA .xml / .xml.p7m / .zip)
ai totali di acquisto per categoria (carne, sfarinati, vino, birra, ...) al netto IVA.

Uso:
  python scripts/fatturapa/parse_fatture.py --in <cartella o zip con le fatture> \
      --out output/acquisti_2025_categorie.xlsx [--anno 2025] [--piva 01241890258] \
      [--rules scripts/fatturapa/regole_categorie.json] [--csv]

Cosa fa:
  1. apre ogni .xml, .xml.p7m (firma CAdES, via openssl o estrazione diretta) e .zip;
  2. legge cedente (fornitore), tipo/data/numero documento, righe di dettaglio e riepilogo IVA;
  3. classifica ogni riga con le regole in regole_categorie.json (parole chiave + fornitore + aliquota);
  4. scrive un Excel con: Totali_ISA, Totali, Fornitori, Righe, Da_verificare, Fatture, Quadratura.

Le note di credito (TD04, TD08) entrano con segno negativo. Le righe accessorie (trasporto, cauzioni,
sconti) restano fuori dalle categorie alimentari. Dipendenze: lxml, openpyxl.
"""
from __future__ import annotations

import argparse
import base64
import csv
import io
import json
import os
import re
import subprocess
import sys
import tempfile
import unicodedata
import zipfile
from collections import defaultdict
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path

from lxml import etree
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FE_BLOCK_RE = re.compile(rb"<(?:[A-Za-z0-9_]+:)?FatturaElettronica[\s>][\s\S]*?</(?:[A-Za-z0-9_]+:)?FatturaElettronica>")
B64_RE = re.compile(rb"^[A-Za-z0-9+/=\r\n\s]+$")
NEGATIVE_TYPES = {"TD04", "TD08"}  # nota di credito, nota di credito semplificata
PRIORITY = [
    "carne", "pesce", "sfarinati", "vino", "birra", "altre_bevande_alcoliche",
    "bevande_analcoliche", "latticini", "ortofrutta", "pane_pasta_riso",
    "altro_alimentare", "non_alimentare",
]
ISA_CATEGORIES = ["carne", "sfarinati", "vino", "birra"]


# ----------------------------------------------------------------------------- utilità

def norm(s: str | None) -> str:
    """minuscolo, senza accenti, spazi singoli; aggiunge uno spazio finale per le radici con spazio."""
    if not s:
        return ""
    s = unicodedata.normalize("NFKD", s)
    s = "".join(c for c in s if not unicodedata.combining(c))
    s = re.sub(r"\s+", " ", s.lower()).strip()
    return s + " "


def dec(s: str | None, default: Decimal = Decimal("0")) -> Decimal:
    if s is None:
        return default
    s = s.strip().replace(",", ".")
    if not s:
        return default
    try:
        return Decimal(s)
    except InvalidOperation:
        return default


def ln(el) -> str:
    return etree.QName(el).localname if isinstance(el.tag, str) else ""


def child(el, *names):
    """naviga per localname (namespace-agnostico); ritorna il primo elemento o None."""
    cur = el
    for n in names:
        nxt = None
        for c in cur:
            if ln(c) == n:
                nxt = c
                break
        if nxt is None:
            return None
        cur = nxt
    return cur


def text(el, *names) -> str:
    c = child(el, *names) if names else el
    return (c.text or "").strip() if c is not None and c.text else ""


def children(el, name):
    return [c for c in el if ln(c) == name]


# ----------------------------------------------------------------------------- lettura file

def unwrap_p7m(data: bytes) -> bytes | None:
    """Estrae l'XML da un .p7m (CAdES DER o base64). Prima openssl, poi estrazione diretta."""
    candidates = [data]
    if B64_RE.match(data[:4000] or b"x"):
        try:
            candidates.append(base64.b64decode(re.sub(rb"\s", b"", data)))
        except Exception:
            pass
    for blob in candidates:
        for inform in ("DER", "PEM"):
            with tempfile.NamedTemporaryFile(suffix=".p7m", delete=False) as tf:
                tf.write(blob)
                tmp = tf.name
            try:
                r = subprocess.run(
                    ["openssl", "cms", "-verify", "-noverify", "-inform", inform, "-in", tmp],
                    capture_output=True, timeout=60,
                )
                if r.returncode == 0 and b"FatturaElettronica" in r.stdout:
                    return r.stdout
                r = subprocess.run(
                    ["openssl", "smime", "-verify", "-noverify", "-inform", inform, "-in", tmp],
                    capture_output=True, timeout=60,
                )
                if r.returncode == 0 and b"FatturaElettronica" in r.stdout:
                    return r.stdout
            except (OSError, subprocess.SubprocessError):
                pass
            finally:
                try:
                    os.unlink(tmp)
                except OSError:
                    pass
        m = FE_BLOCK_RE.search(blob)
        if m:
            return m.group(0)
    return None


def iter_sources(path: Path):
    """Genera (nome, bytes) per ogni fattura trovata in file/cartelle/zip (ricorsivo)."""
    if path.is_dir():
        for p in sorted(path.rglob("*")):
            if p.is_file():
                yield from iter_sources(p)
        return
    name = path.name
    low = name.lower()
    if low.endswith(".zip"):
        try:
            with zipfile.ZipFile(path) as z:
                for info in z.infolist():
                    if info.is_dir():
                        continue
                    inner = info.filename.lower()
                    data = z.read(info)
                    if inner.endswith(".zip"):
                        with tempfile.NamedTemporaryFile(suffix=".zip", delete=False) as tf:
                            tf.write(data)
                            tmp = Path(tf.name)
                        try:
                            yield from iter_sources(tmp)
                        finally:
                            tmp.unlink(missing_ok=True)
                    elif inner.endswith((".xml", ".p7m")):
                        yield f"{name}/{info.filename}", data
        except zipfile.BadZipFile:
            print(f"[avviso] zip non leggibile: {path}", file=sys.stderr)
        return
    if low.endswith((".xml", ".p7m")):
        yield name, path.read_bytes()


def parse_xml(data: bytes):
    if b"FatturaElettronica" not in data:
        unwrapped = unwrap_p7m(data)
        if unwrapped is None:
            return None
        data = unwrapped
    if data.lstrip().startswith(b"<?xml") is False and not data.lstrip().startswith(b"<"):
        m = FE_BLOCK_RE.search(data)
        if m:
            data = m.group(0)
    parser = etree.XMLParser(recover=True, huge_tree=True, remove_blank_text=True)
    try:
        root = etree.fromstring(data, parser=parser)
    except etree.XMLSyntaxError:
        return None
    if root is None:
        return None
    # alcuni file .p7m contengono l'XML dentro un blocco: cerchiamo FatturaElettronica ovunque
    if ln(root) != "FatturaElettronica":
        for el in root.iter():
            if ln(el) == "FatturaElettronica":
                root = el
                break
    return root


# ----------------------------------------------------------------------------- modello

@dataclass
class Riga:
    file: str
    fornitore: str
    fornitore_piva: str
    tipo_doc: str
    data_doc: str
    numero_doc: str
    n_linea: str
    descrizione: str
    quantita: Decimal
    um: str
    prezzo_unitario: Decimal
    imponibile: Decimal
    aliquota: Decimal
    natura: str
    categoria: str = ""
    sottocategoria: str = ""
    motivo: str = ""
    da_verificare: bool = False
    accessoria: bool = False


@dataclass
class Fattura:
    file: str
    fornitore: str
    fornitore_piva: str
    tipo_doc: str
    data_doc: str
    numero_doc: str
    cessionario_piva: str
    imponibile_riepilogo: Decimal
    imposta_riepilogo: Decimal
    totale_documento: Decimal
    righe: list = field(default_factory=list)


def read_fattura(name: str, root) -> list[Fattura]:
    out = []
    header = child(root, "FatturaElettronicaHeader")
    bodies = children(root, "FatturaElettronicaBody")
    if header is None or not bodies:
        return out
    ced = child(header, "CedentePrestatore", "DatiAnagrafici")
    forn_piva = ""
    forn = ""
    if ced is not None:
        forn_piva = text(ced, "IdFiscaleIVA", "IdCodice") or text(ced, "CodiceFiscale")
        anag = child(ced, "Anagrafica")
        if anag is not None:
            forn = text(anag, "Denominazione") or (text(anag, "Nome") + " " + text(anag, "Cognome")).strip()
    ces = child(header, "CessionarioCommittente", "DatiAnagrafici")
    ces_piva = ""
    if ces is not None:
        ces_piva = text(ces, "IdFiscaleIVA", "IdCodice") or text(ces, "CodiceFiscale")

    for body in bodies:
        dg = child(body, "DatiGenerali", "DatiGeneraliDocumento")
        tipo = text(dg, "TipoDocumento") if dg is not None else ""
        data_doc = text(dg, "Data") if dg is not None else ""
        numero = text(dg, "Numero") if dg is not None else ""
        totale = dec(text(dg, "ImportoTotaleDocumento")) if dg is not None else Decimal("0")
        segno = Decimal("-1") if tipo in NEGATIVE_TYPES else Decimal("1")
        f = Fattura(name, forn, forn_piva, tipo, data_doc, numero, ces_piva, Decimal("0"), Decimal("0"), totale * segno)
        dbs = child(body, "DatiBeniServizi")
        if dbs is not None:
            for r in children(dbs, "DatiRiepilogo"):
                f.imponibile_riepilogo += dec(text(r, "ImponibileImporto")) * segno
                f.imposta_riepilogo += dec(text(r, "Imposta")) * segno
            for l in children(dbs, "DettaglioLinee"):
                q = dec(text(l, "Quantita"), Decimal("1"))
                f.righe.append(Riga(
                    file=name, fornitore=forn, fornitore_piva=forn_piva, tipo_doc=tipo,
                    data_doc=data_doc, numero_doc=numero, n_linea=text(l, "NumeroLinea"),
                    descrizione=text(l, "Descrizione"), quantita=q, um=text(l, "UnitaMisura"),
                    prezzo_unitario=dec(text(l, "PrezzoUnitario")),
                    imponibile=dec(text(l, "PrezzoTotale")) * segno,
                    aliquota=dec(text(l, "AliquotaIVA")), natura=text(l, "Natura"),
                ))
        out.append(f)
    return out


# ----------------------------------------------------------------------------- classificazione

class Classifier:
    def __init__(self, rules: dict):
        self.rules = rules
        self.cats = rules["categorie"]
        self.compiled = {}
        for cat, spec in self.cats.items():
            self.compiled[cat] = (
                [self._rx(k) for k in spec.get("keywords", [])],
                [self._rx(k) for k in spec.get("escludi", [])],
                {sub: [self._rx(k) for k in kws] for sub, kws in spec.get("sotto", {}).items()},
                {Decimal(str(a)) for a in spec.get("aliquote_attese", [])},
            )
        self.accessorie = [self._rx(k) for k in rules.get("accessorie", [])]
        self.fornitori = [(norm(f["match"]).strip(), f["default"], f.get("nota", "")) for f in rules.get("fornitori", [])]

    @staticmethod
    def _rx(stem: str):
        stem_n = norm(stem).rstrip() + (" " if stem.endswith(" ") else "")
        return re.compile(r"(?<![a-z0-9])" + re.escape(stem_n))

    def classify(self, r: Riga):
        d = norm(r.descrizione)
        forn = norm(r.fornitore)
        # accessorie: fuori dalle categorie
        if any(rx.search(d) for rx in self.accessorie) and not any(
            rx.search(d) for rx in self.compiled["carne"][0] + self.compiled["vino"][0] + self.compiled["birra"][0] + self.compiled["sfarinati"][0]
        ):
            r.categoria, r.accessoria, r.motivo = "accessoria", True, "riga accessoria (trasporto/sconto/cauzione)"
            return
        scores = {}
        for cat, (kws, excl, subs, _) in self.compiled.items():
            if any(rx.search(d) for rx in excl):
                continue
            hits = [rx.pattern for rx in kws if rx.search(d)]
            if hits:
                scores[cat] = hits
        chosen, motivo, verifica = "", "", False
        if scores:
            best = max(len(h) for h in scores.values())
            top = [c for c, h in scores.items() if len(h) == best]
            top.sort(key=lambda c: PRIORITY.index(c) if c in PRIORITY else 99)
            chosen = top[0]
            motivo = "parole chiave: " + ", ".join(re.sub(r"\\", "", h.split(")", 1)[-1]) for h in scores[chosen][:3])
            food_conflict = [c for c in scores if c in ISA_CATEGORIES and c != chosen]
            if len(top) > 1 or food_conflict:
                verifica = True
                motivo += " | ambigua con " + ", ".join(c for c in scores if c != chosen)
        else:
            for match, default, nota in self.fornitori:
                if match and match in forn:
                    chosen, motivo, verifica = default, f"default fornitore ({nota})", True
                    break
            if not chosen:
                chosen, motivo, verifica = "altro", "nessuna parola chiave", True
        # sottocategoria
        subs = self.compiled.get(chosen, (None, None, {}, set()))[2] if chosen in self.compiled else {}
        for sub, rxs in subs.items():
            if any(rx.search(d) for rx in rxs):
                r.sottocategoria = sub
                break
        # controllo aliquota
        expected = self.compiled[chosen][3] if chosen in self.compiled else set()
        if expected and r.aliquota not in expected and not r.natura:
            verifica = True
            motivo += f" | aliquota {r.aliquota}% inattesa per {chosen} (attese {sorted(int(a) for a in expected)})"
        r.categoria, r.motivo, r.da_verificare = chosen, motivo, verifica


# ----------------------------------------------------------------------------- output

HEAD_FILL = PatternFill("solid", fgColor="DDEBF7")
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")


def _sheet(wb, title, header, rows, widths=None, warn_col=None):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for row in rows:
        ws.append([float(v) if isinstance(v, Decimal) else v for v in row])
    if warn_col is not None:
        for row in ws.iter_rows(min_row=2):
            if row[warn_col].value in (True, "SI", "sì", "SÌ"):
                for c in row:
                    c.fill = WARN_FILL
    for i, w in enumerate(widths or [], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    return ws


def write_xlsx(out: Path, fatture: list[Fattura], anno: str | None):
    righe = [r for f in fatture for r in f.righe]
    wb = Workbook()
    wb.remove(wb.active)

    # Totali ISA (le 4 categorie richieste) con split confermato / da verificare
    tot = defaultdict(lambda: {"ok": Decimal("0"), "ver": Decimal("0"), "n": 0, "fatt": set()})
    for r in righe:
        t = tot[r.categoria]
        t["ver" if r.da_verificare else "ok"] += r.imponibile
        t["n"] += 1
        t["fatt"].add((r.fornitore_piva, r.numero_doc, r.data_doc))
    isa_rows = []
    for c in ISA_CATEGORIES:
        t = tot.get(c, {"ok": Decimal("0"), "ver": Decimal("0"), "n": 0, "fatt": set()})
        isa_rows.append([c, t["ok"], t["ver"], t["ok"] + t["ver"], t["n"], len(t["fatt"])])
    _sheet(wb, "Totali_ISA",
           ["Categoria", "Imponibile confermato €", "Imponibile da verificare €", "Totale €", "N. righe", "N. fatture"],
           isa_rows, [22, 24, 26, 16, 10, 12])

    all_rows = []
    for c in sorted(tot, key=lambda c: (PRIORITY.index(c) if c in PRIORITY else 99, c)):
        t = tot[c]
        all_rows.append([c, t["ok"], t["ver"], t["ok"] + t["ver"], t["n"], len(t["fatt"])])
    all_rows.append(["TOTALE", sum(t["ok"] for t in tot.values()), sum(t["ver"] for t in tot.values()),
                     sum(t["ok"] + t["ver"] for t in tot.values()), len(righe), len({(f.fornitore_piva, f.numero_doc, f.data_doc) for f in fatture})])
    _sheet(wb, "Totali", ["Categoria", "Imponibile confermato €", "Imponibile da verificare €", "Totale €", "N. righe", "N. fatture"],
           all_rows, [26, 24, 26, 16, 10, 12])

    # fornitori x categoria
    fx = defaultdict(lambda: defaultdict(Decimal))
    for r in righe:
        fx[(r.fornitore, r.fornitore_piva)][r.categoria] += r.imponibile
    cats = [c for c in PRIORITY if any(c in d for d in fx.values())] + sorted(
        {c for d in fx.values() for c in d if c not in PRIORITY})
    frows = []
    for (forn, piva), d in sorted(fx.items(), key=lambda kv: -sum(kv[1].values())):
        frows.append([forn, piva] + [d.get(c, Decimal("0")) for c in cats] + [sum(d.values())])
    _sheet(wb, "Fornitori", ["Fornitore", "P.IVA"] + cats + ["Totale €"], frows, [34, 14] + [14] * (len(cats) + 1))

    rh = ["File", "Fornitore", "P.IVA", "Tipo", "Data", "Numero", "Linea", "Descrizione", "Q.tà", "UM",
          "Prezzo unit.", "Imponibile €", "IVA %", "Natura", "Categoria", "Sottocategoria", "Da verificare", "Motivo"]
    rrows = [[r.file, r.fornitore, r.fornitore_piva, r.tipo_doc, r.data_doc, r.numero_doc, r.n_linea, r.descrizione,
              r.quantita, r.um, r.prezzo_unitario, r.imponibile, r.aliquota, r.natura, r.categoria, r.sottocategoria,
              "SI" if r.da_verificare else "", r.motivo] for r in righe]
    widths = [28, 28, 13, 6, 11, 12, 6, 60, 8, 5, 11, 13, 7, 7, 18, 14, 11, 50]
    _sheet(wb, "Righe", rh, rrows, widths, warn_col=16)
    _sheet(wb, "Da_verificare", rh, [row for row, r in zip(rrows, righe) if r.da_verificare], widths, warn_col=16)

    fh = ["File", "Fornitore", "P.IVA", "Tipo", "Data", "Numero", "Cessionario P.IVA", "Imponibile riepilogo €",
          "Imposta €", "Totale documento €", "N. righe", "Somma righe €", "Differenza €"]
    frows2, quad = [], []
    for f in fatture:
        somma = sum((r.imponibile for r in f.righe), Decimal("0"))
        diff = f.imponibile_riepilogo - somma
        frows2.append([f.file, f.fornitore, f.fornitore_piva, f.tipo_doc, f.data_doc, f.numero_doc, f.cessionario_piva,
                       f.imponibile_riepilogo, f.imposta_riepilogo, f.totale_documento, len(f.righe), somma, diff])
        if abs(diff) > Decimal("0.05"):
            quad.append(frows2[-1])
    _sheet(wb, "Fatture", fh, frows2, [28, 28, 13, 6, 11, 12, 15, 20, 12, 18, 9, 15, 13])
    _sheet(wb, "Quadratura", fh, quad, [28, 28, 13, 6, 11, 12, 15, 20, 12, 18, 9, 15, 13])

    note = wb.create_sheet("Note", 0)
    note["A1"] = "Acquisti per categoria da fatture elettroniche" + (f" – anno {anno}" if anno else "")
    note["A1"].font = Font(bold=True, size=13)
    lines = [
        "Importi al netto IVA (imponibile di riga, PrezzoTotale). Note di credito TD04/TD08 con segno negativo.",
        "Totali_ISA: le 4 categorie chieste dallo studio, con separazione tra righe confermate e righe da verificare.",
        "Da_verificare: righe ambigue (più categorie, aliquota inattesa, default fornitore, nessuna parola chiave). Correggere la colonna Categoria e riaggiornare i totali.",
        "Quadratura: fatture in cui la somma delle righe non coincide con l'imponibile del riepilogo IVA (sconti di piede, arrotondamenti, righe descrittive).",
        "Regole di classificazione: scripts/fatturapa/regole_categorie.json (modificabili).",
    ]
    for i, l in enumerate(lines, start=3):
        note[f"A{i}"] = l
    note.column_dimensions["A"].width = 120
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)


def write_csv(out: Path, fatture: list[Fattura]):
    p = out.with_suffix(".righe.csv")
    with p.open("w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh, delimiter=";")
        w.writerow(["file", "fornitore", "piva", "tipo", "data", "numero", "linea", "descrizione", "quantita", "um",
                    "prezzo_unitario", "imponibile", "aliquota", "natura", "categoria", "sottocategoria", "da_verificare", "motivo"])
        for f in fatture:
            for r in f.righe:
                w.writerow([r.file, r.fornitore, r.fornitore_piva, r.tipo_doc, r.data_doc, r.numero_doc, r.n_linea,
                            r.descrizione, r.quantita, r.um, r.prezzo_unitario, r.imponibile, r.aliquota, r.natura,
                            r.categoria, r.sottocategoria, "SI" if r.da_verificare else "", r.motivo])
    return p


# ----------------------------------------------------------------------------- main

def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--in", dest="src", required=True, help="cartella, file o zip con le fatture")
    ap.add_argument("--out", required=True, help="file Excel di output")
    ap.add_argument("--anno", help="tieni solo le fatture con Data che inizia con questo anno (es. 2025)")
    ap.add_argument("--piva", help="tieni solo le fatture intestate a questa P.IVA (cessionario)")
    ap.add_argument("--rules", default=str(Path(__file__).with_name("regole_categorie.json")))
    ap.add_argument("--csv", action="store_true", help="scrivi anche il CSV delle righe")
    args = ap.parse_args(argv)

    rules = json.loads(Path(args.rules).read_text(encoding="utf-8"))
    clf = Classifier(rules)
    src = Path(args.src)
    if not src.exists():
        print(f"[errore] percorso inesistente: {src}", file=sys.stderr)
        return 2

    fatture: list[Fattura] = []
    n_files = n_bad = 0
    for name, data in iter_sources(src):
        n_files += 1
        root = parse_xml(data)
        if root is None:
            n_bad += 1
            print(f"[avviso] non leggibile: {name}", file=sys.stderr)
            continue
        for f in read_fattura(name, root):
            if args.anno and not f.data_doc.startswith(args.anno):
                continue
            if args.piva and f.cessionario_piva and f.cessionario_piva != args.piva:
                continue
            for r in f.righe:
                clf.classify(r)
            fatture.append(f)

    out = Path(args.out)
    write_xlsx(out, fatture, args.anno)
    if args.csv:
        write_csv(out, fatture)

    righe = [r for f in fatture for r in f.righe]
    print(f"file letti: {n_files} (non leggibili: {n_bad}); fatture: {len(fatture)}; righe: {len(righe)}")
    tot = defaultdict(Decimal)
    ver = defaultdict(Decimal)
    for r in righe:
        tot[r.categoria] += r.imponibile
        if r.da_verificare:
            ver[r.categoria] += r.imponibile
    for c in ISA_CATEGORIES:
        print(f"  {c:<12} {tot[c]:>12.2f} €  (da verificare {ver[c]:.2f} €)")
    print(f"output: {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
